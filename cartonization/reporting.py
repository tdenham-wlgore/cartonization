"""Portable Excel/JSON snapshots using the package's existing openpyxl dependency."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def write_json(path, value):
    """Write strict JSON (no NaN/Infinity), preserving text identifiers."""
    Path(path).write_text(
        json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8"
    )


def _value(cell, value):
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"  # Treat identifiers/notes as text, never executable formulas.


def _table(wb, name, headers, rows, *, title=None, formats=None, widths=None):
    sheet = wb.create_sheet(name)
    sheet.sheet_view.showGridLines = False
    header_row = 3 if title else 1
    if title:
        _value(sheet.cell(1, 1), title)
        sheet.cell(1, 1).font = Font(name="Arial", size=14, bold=True, color="243746")
    for j, h in enumerate(headers, 1):
        cell = sheet.cell(header_row, j)
        _value(cell, h)
        cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="34495E")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[header_row].height = 34
    count = 0
    for row_num, values in enumerate(rows, header_row + 1):
        if row_num > 1048576:
            raise ValueError(f"{name}: too many rows for Excel; reduce the output scope.")
        count += 1
        for j, value in enumerate(values, 1):
            cell = sheet.cell(row_num, j)
            _value(cell, value)
            cell.font = Font(name="Arial", size=11, color="243746")
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=isinstance(value, str),
                horizontal="left" if isinstance(value, str) else "right",
            )
            if formats and j in formats:
                cell.number_format = formats[j]
            if row_num % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F2F5F7")
        sheet.row_dimensions[row_num].height = (
            30 if any(isinstance(x, str) and len(x) > 50 for x in values) else 21
        )
    for j in range(1, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(j)].width = (widths or {}).get(j, 22)
    if count:
        sheet.auto_filter.ref = (
            f"A{header_row}:{get_column_letter(len(headers))}{header_row + count}"
        )
    sheet.freeze_panes = f"B{header_row + 1}"
    sheet.print_options.horizontalCentered = True
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    return sheet


def _workbook():
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def export_issue_report(issues, directory):
    """Export actionable issues, without successful-looking analysis totals."""
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    write_json(directory / "issues.json", {"issues": issues})
    wb = _workbook()
    _table(
        wb,
        "Issues",
        ["Location", "Problem", "Affected demand"],
        [
            [
                i.get("location", ""),
                i.get("message", ""),
                json.dumps(
                    {k: v for k, v in i.items() if k not in ("location", "message")},
                    ensure_ascii=False,
                )
                if len(i) > 2
                else "",
            ]
            for i in issues
        ],
        title="Input or analysis issues",
        widths={1: 35, 2: 95, 3: 45},
    )
    wb.save(directory / "issues.xlsx")


def export_report(results, directory, comparison=None, usage_comparison=None):
    """Write report.xlsx and results.json. Values are reproducible model snapshots."""
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    comparison, usage_comparison = comparison or [], usage_comparison or []
    wb = _workbook()
    money, percent = '"$"#,##0.00', "0.0%"
    _table(
        wb,
        "Summary",
        [
            "Scenario",
            "Modeled cost",
            "Shippers",
            "Packing efficiency",
            "Demand coverage",
            "Included orders",
        ],
        [
            [
                r.name,
                r.total_cost,
                sum(r.total_count.values()),
                r.packing_efficiency / 100,
                r.coverage["included_fraction"],
                r.coverage["included_frequency"],
            ]
            for r in results
        ],
        title="Cartonization results",
        formats={2: money, 3: "#,##0", 4: percent, 5: percent, 6: "#,##0"},
        widths={1: 28, 2: 22, 3: 15, 4: 22, 5: 22, 6: 21},
    )
    if comparison:
        _table(
            wb,
            "Comparison",
            [
                "Scenario",
                "Baseline",
                "Cost before",
                "Cost after",
                "Cost change",
                "Efficiency before",
                "Efficiency after",
                "Efficiency change (pp)",
            ],
            [
                [
                    r[k]
                    for k in (
                        "scenario",
                        "baseline",
                        "cost_before",
                        "cost_after",
                        "cost_change",
                        "efficiency_before",
                        "efficiency_after",
                        "efficiency_change_pp",
                    )
                ]
                for r in comparison
            ],
            title="Before and after",
            formats={3: money, 4: money, 5: money, 6: percent, 7: percent, 8: "0.0"},
            widths={1: 25, 2: 25, 3: 19, 4: 19, 5: 19, 6: 21, 7: 21, 8: 24},
        )
        _table(
            wb,
            "Usage Changes",
            ["Scenario", "Baseline", "Shipper", "Before", "After", "Change"],
            [
                [r[k] for k in ("scenario", "baseline", "shipper", "before", "after", "change")]
                for r in usage_comparison
            ],
            title="Shipper usage changes",
            formats={4: "#,##0", 5: "#,##0", 6: "#,##0"},
        )
    _table(
        wb,
        "Shipper Usage",
        ["Scenario", "Shipper", "Quantity", "Modeled cost per shipper", "Modeled total cost"],
        [
            [r.name, sid, qty, r.shippers[sid]["cost"], qty * r.shippers[sid]["cost"]]
            for r in results
            for sid, qty in r.total_count.items()
        ],
        title="Shipper usage",
        formats={3: "#,##0", 4: money, 5: money},
        widths={4: 28, 5: 25},
    )
    _table(
        wb,
        "Max Fit",
        ["Scenario", "Shipper", "Carton", "Four-quadrant capacity"],
        [[r.name, sid, cid, fit] for r in results for (sid, cid), fit in r.max_fits.items()],
        title="Single-carton capacity",
        formats={4: "#,##0"},
        widths={4: 28},
    )
    _table(
        wb,
        "Packings",
        ["Scenario", "Packing", "Shipper", "Capacity method", "Carton capacities"],
        [
            [
                r.name,
                p["packing_id"],
                p["shipper_id"],
                p["fit_method"],
                ", ".join(f"{cid}: {qty}" for cid, qty in p["counts"].items() if qty),
            ]
            for r in results
            for p in r.packings
        ],
        title="Packing capacities",
        widths={2: 26, 4: 25, 5: 72},
    )
    _table(
        wb,
        "Profile Results",
        [
            "Scenario",
            "Profile",
            "Frequency",
            "Cost per order",
            "Solver result",
            "Cartons",
            "Selected shippers",
        ],
        [
            [
                r.name,
                p["profile"],
                p["frequency"],
                p["cost_per_order"],
                p["status"],
                ", ".join(f"{c}: {q}" for c, q in p["counts"].items() if q),
                ", ".join(f"{s['shipper_id']}: {s['quantity']}" for s in p["selections"]),
            ]
            for r in results
            for p in r.profiles
        ],
        title="Shipment profile results",
        formats={3: "#,##0", 4: money},
        widths={1: 25, 2: 12, 3: 16, 4: 20, 5: 35, 6: 50, 7: 50},
    )
    metadata = []
    for r in results:
        for section, values in (("Settings", r.settings), ("Coverage", r.coverage)):
            metadata.extend(
                [r.name, section, k, json.dumps(v, ensure_ascii=False)] for k, v in values.items()
            )
        for f in r.provenance.get("input_files", []):
            metadata.append([r.name, "Input", f["name"], f["sha256"]])
        metadata.append(
            [r.name, "Source sheet", "Shipment history", r.provenance.get("history_sheet", "")]
        )
        metadata.extend(
            [r.name, "Software", k, v] for k, v in r.provenance.get("dependencies", {}).items()
        )
        metadata.append([r.name, "Software", "Python", r.provenance.get("python", "")])
        metadata.extend([r.name, "Notice", "", w] for w in r.warnings)
    _table(
        wb,
        "Settings and Sources",
        ["Scenario", "Section", "Setting / source", "Value"],
        metadata,
        title="Settings and sources",
        widths={1: 28, 2: 22, 3: 42, 4: 90},
    )
    notes = [
        [
            "Snapshot",
            "These are calculated outputs. Change saved scenario settings and rerun to update them.",
        ],
        [
            "Default objective",
            "volume_penalty minimizes rounded shipper volume plus the largest candidate shipper volume per shipper. It does not guarantee minimum shipping cost or the fewest shippers.",
        ],
        [
            "Cost option",
            "shipping_cost minimizes supplied modeled costs. The built-in Zone 4 formula is a legacy dimensional-weight estimate, not a current carrier tariff.",
        ],
        [
            "Capacity interpolation",
            "By default, mixed packings involving any carton with single-type capacity >= 10 use a normalized capacity estimate. No expensive 3D search is required in those cases.",
        ],
        [
            "Geometry",
            "Four-quadrant and 3D checks find arrangements heuristically. Capacity limits, packing caps and interpolation can affect the solution. No method here certifies global packing optimality or physical pack-test approval.",
        ],
        [
            "Constraints",
            "Dimensions must use inches. Shipper dimensions are usable interior dimensions; carton dimensions are outer dimensions. Weight, inserts, fragility, orientation restrictions and supplier availability are not modeled.",
        ],
        [
            "Scope",
            "Only selected cartons contribute to the analyzed demand. Container groups can overlap orders; do not sum order frequencies across groups as a count of unique orders.",
        ],
        [
            "Coverage",
            "Sampled/filtered results are observed totals for included demand; no extrapolation is applied. Zero-frequency rows and orders containing no selected cartons are excluded before scope totals.",
        ],
        [
            "Change signs",
            "Changes are after minus before. Negative cost changes indicate savings. Efficiency changes are percentage points.",
        ],
        [
            "Reproducibility",
            "Input hashes, settings and software versions are recorded. Equal optima can select different shippers across solver versions/platforms.",
        ],
    ]
    _table(
        wb,
        "Notes",
        ["Topic", "Explanation"],
        notes,
        title="Reading the results",
        widths={1: 28, 2: 115},
    )
    for row in range(4, 4 + len(notes)):
        wb["Notes"].row_dimensions[row].height = 48
    wb.save(directory / "report.xlsx")
    write_json(
        directory / "results.json",
        {
            "scenarios": [r.to_dict() for r in results],
            "comparisons": comparison,
            "usage_changes": usage_comparison,
        },
    )
    return directory / "report.xlsx"


def export_prepared(prepared, directory):
    """Write solver-ready profiles, retaining source sheet names and header row 1."""
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    wb = _workbook()
    for name, profiles in prepared["sheets"].items():
        _table(
            wb,
            name,
            prepared["cartons"] + ["FREQUENCY"],
            [[p["counts"][c] for c in prepared["cartons"]] + [p["frequency"]] for p in profiles],
        )
    info_name = "Preparation Statistics"
    while info_name.lower() in [s.lower() for s in wb.sheetnames]:
        info_name = "_" + info_name
    fields = list(prepared["statistics"][0])
    _table(wb, info_name, fields, [[s[k] for k in fields] for s in prepared["statistics"]])
    wb.save(directory / "shipment_profiles.xlsx")
    write_json(directory / "preparation.json", prepared)
    return directory / "shipment_profiles.xlsx"
