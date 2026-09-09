"""Streaming Excel readers; identifiers stay text and invalid counts never truncate."""

from __future__ import annotations

from collections import Counter

from openpyxl import load_workbook

from .validation import InputError, identifier, identifiers, integer, invalid


def read_xlsx_sheet_rows(filename, sheet_name):
    """Return worksheet rows, including the header; prefer streaming for large files."""
    wb = load_workbook(filename, read_only=True, data_only=True)
    try:
        return [list(row) for row in wb[sheet_name].iter_rows(values_only=True)]
    finally:
        wb.close()


def iter_sheet_records(filename, sheet_name, usecols=None):
    """Yield (Excel row number, dict) without materializing the workbook.

    Inputs must contain values, not formulas, to avoid relying on missing/stale
    cached results. Blank rows are ignored; nonblank headers must be unique.
    """
    wb = load_workbook(filename, read_only=True, data_only=False)
    try:
        if sheet_name not in wb.sheetnames:
            raise invalid(
                str(filename), f"Sheet {sheet_name!r} not found; available: {wb.sheetnames}."
            )
        rows = wb[sheet_name].iter_rows(values_only=True)
        first = next(rows, ())
        headers = [str(v).strip() if v is not None else "" for v in first]
        duplicates = [k for k, v in Counter(h for h in headers if h).items() if v > 1]
        if duplicates:
            raise invalid(f"{sheet_name}!1", f"Duplicate headers: {duplicates}.")
        positions = {h: i for i, h in enumerate(headers) if h}
        keep = list(usecols) if usecols is not None else list(positions)
        missing = [c for c in keep if c not in positions]
        if missing:
            raise invalid(f"{sheet_name}!1", f"Missing columns: {missing}.")
        if not positions:
            raise invalid(f"{sheet_name}!1", "A header row is required.")
        for row_num, values in enumerate(rows, 2):
            if all(v is None for v in values):
                continue
            row = {c: values[positions[c]] if positions[c] < len(values) else None for c in keep}
            for col, value in row.items():
                if isinstance(value, str) and value.startswith("="):
                    raise invalid(
                        f"{sheet_name}!{col}{row_num}",
                        "Paste formula results as values before analysis.",
                    )
            yield row_num, row
    finally:
        wb.close()


def stream_xlsx_as_dicts(filename, sheet_name, usecols=None):
    """Yield dictionaries from a sheet. Consume fully or close the generator."""
    for _, row in iter_sheet_records(filename, sheet_name, usecols):
        yield row


def read_xlsx_as_dicts(filename, sheet_name, usecols=None):
    """Compatibility list wrapper around the streaming reader."""
    return list(stream_xlsx_as_dicts(filename, sheet_name, usecols))


def load_reference_mapping(ref_file, sheet_name, usecols, id_col, allowed_ids, build_value):
    """Load selected IDs, rejecting missing/duplicate selected records."""
    allowed = set(identifiers(allowed_ids, "allowed_ids"))
    out = {}
    for row_num, row in iter_sheet_records(ref_file, sheet_name, usecols):
        rid = identifier(row[id_col], f"{sheet_name}!{id_col}{row_num}")
        if rid not in allowed:
            continue
        if rid in out:
            raise invalid(f"{sheet_name}!{id_col}{row_num}", f"Duplicate ID {rid!r}.")
        try:
            out[rid] = build_value(row)
        except (ValueError, TypeError, OverflowError) as exc:
            raise invalid(f"{sheet_name}!row {row_num} ({rid})", str(exc)) from exc
    missing = sorted(allowed - out.keys())
    if missing:
        raise invalid(sheet_name, f"Selected IDs missing from reference: {missing}.")
    return out


def build_shipment_profiles(history_file, sheet_name, carton_list, frequency_col="FREQUENCY"):
    """Aggregate equal carton-count vectors and their frequencies.

    Blank counts/frequencies are zero. Negative/fractional values are errors.
    Zero-frequency rows and orders with no selected cartons contribute no demand.
    """
    carton_list = identifiers(carton_list, "carton_list")
    if frequency_col in carton_list:
        raise invalid("frequency_col", "The frequency column cannot also be a carton ID.")
    agg = Counter()
    issues = []
    for row_num, row in iter_sheet_records(history_file, sheet_name, [frequency_col] + carton_list):
        try:
            freq = integer(
                row[frequency_col], f"{sheet_name}!{frequency_col}{row_num}", blank_zero=True
            )
            counts = tuple(
                integer(row[c], f"{sheet_name}!{c}{row_num}", blank_zero=True) for c in carton_list
            )
            if freq and any(counts):
                agg[counts] += freq
        except InputError as exc:
            issues.extend(exc.issues)
    if issues:
        raise InputError(f"Shipment history has {len(issues)} invalid row(s).", issues)
    return [
        {"frequency": f, "counts": dict(zip(carton_list, counts))}
        for counts, f in sorted(agg.items())
    ]
