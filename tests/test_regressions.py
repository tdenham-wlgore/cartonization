"""Regression tests for model correctness and the portable handoff workflows."""

import itertools
import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from cartonization.analysis import analyze_loaded, calculate_cost_count
from cartonization.cli import main
from cartonization.costs import shipping_cost_zone4
from cartonization.io import build_shipment_profiles, read_xlsx_as_dicts
from cartonization.packing import (
    _select_spread_packings,
    comingle_frontier_points,
    comingle_tester,
    generate_packings,
    max_cartons_four_quadrant,
)
from cartonization.reporting import export_report
from cartonization.solver import solve_min_integer, upper_bounds_for_solve_min_integer
from cartonization.validation import AnalysisError, InputError
from cartonization.workflows import (
    load_config,
    prepare_shipment_history,
    run_scenario,
    run_scenarios,
)

ROOT = Path(__file__).resolve().parents[1]


def workbook(path, sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    wb.save(path)


def shipper(volume=1, cost=1):
    return {
        "length": volume,
        "width": 1,
        "height": 1,
        "max_units": 999,
        "volume": volume,
        "cost": cost,
    }


def test_infeasible_solution_must_not_be_returned_as_feasible():
    with pytest.raises(RuntimeError, match="Infeasible"):
        solve_min_integer([1], [[1]], [1], upper_bounds=[0.5])


def test_zero_cost_is_zero_not_infinity():
    x, obj, status = solve_min_integer([1], [[1]], [0], upper_bounds=[1])
    assert (x, obj, status) == ([1], 0, "Optimal")


@pytest.mark.parametrize("b,a", [([], []), ([1], [[]]), ([1, 2], [[1]])])
def test_bad_bounds_input(b, a):
    with pytest.raises(ValueError):
        upper_bounds_for_solve_min_integer(b, a)


def test_covering_solver_matches_small_exhaustive_search():
    for b in ([1, 1], [3, 2], [4, 5]):
        a, costs = [[2, 0, 1], [0, 3, 1]], [4, 3, 2]
        bounds = upper_bounds_for_solve_min_integer(b, a)
        candidates = [
            x
            for x in itertools.product(*(range(n + 1) for n in bounds))
            if all(sum(v * c for v, c in zip(x, row)) >= demand for row, demand in zip(a, b))
        ]
        expected = min(sum(c * v for c, v in zip(costs, x)) for x in candidates)
        _, obj, status = solve_min_integer(b, a, costs, upper_bounds=bounds)
        assert status == "Optimal"
        assert obj == expected


def test_feasible_incumbent_status_and_objective(monkeypatch):
    import pulp

    def fake_solve(prob, solver):
        prob.status = pulp.LpStatusOptimal
        prob.sol_status = pulp.LpSolutionIntegerFeasible
        for v in prob.variables():
            v.varValue = 2
        return prob.status

    monkeypatch.setattr(pulp.LpProblem, "solve", fake_solve)
    assert solve_min_integer([1], [[1]], [3], upper_bounds=[3]) == (
        [2],
        6,
        "Feasible (not proven optimal)",
    )


@pytest.mark.parametrize("value", [1.4, -1, float("nan")])
def test_invalid_incumbents_rejected(monkeypatch, value):
    import pulp

    def fake_solve(prob, solver):
        prob.status = pulp.LpStatusNotSolved
        for v in prob.variables():
            v.varValue = value
        return prob.status

    monkeypatch.setattr(pulp.LpProblem, "solve", fake_solve)
    with pytest.raises(RuntimeError):
        solve_min_integer([1], [[1]], [1], upper_bounds=[1])


def test_geometric_checks_do_not_round_oversized_cartons_down():
    assert not comingle_tester([("C", 1.004, 0.5, 0.5, 1)], ("S", 1, 1, 1, 1), attempts=1)


def test_decimal_grid_boundary_is_exact():
    assert max_cartons_four_quadrant((0.3, 0.3, 0.3), (0.1, 0.1, 0.1)) == 27


def test_geometric_helper_has_no_accidental_weight_capacity():
    assert comingle_tester([("C", 0.5, 0.5, 0.5, 2)], ("S", 1, 0.5, 0.5, 0.25), attempts=2)


def test_default_interpolation_avoids_3d_search_and_rejects_excess_volume(monkeypatch):
    def forbidden(*args, **kwargs):
        raise AssertionError("High-capacity interpolation must not invoke expensive 3D checks.")

    monkeypatch.setattr("cartonization.packing.comingle_tester", forbidden)
    points = comingle_frontier_points(
        ("S", 10, 1, 1, 999, 10), [("A", 1, 1, 1, 10), ("B", 1.1, 1, 1, 9)], [[1, 1]]
    )
    assert any(p.get("A", 0) and p.get("B", 0) for p in points)
    assert all(p.get("A", 0) + 1.1 * p.get("B", 0) <= 10 + 1e-9 for p in points)


def test_low_capacity_and_opt_in_geometry_use_3d_checker(monkeypatch):
    calls = []

    def checker(*args, **kwargs):
        calls.append(kwargs)
        return True

    monkeypatch.setattr("cartonization.packing.comingle_tester", checker)
    comingle_frontier_points(
        ("S", 2, 1, 1, 99, 2), [("A", 1, 1, 1, 2), ("B", 1, 1, 1, 2)], [[1, 1]]
    )
    assert calls
    calls.clear()
    comingle_frontier_points(
        ("S", 10, 1, 1, 99, 10),
        [("A", 1, 1, 1, 10), ("B", 1, 1, 1, 10)],
        [[1, 1]],
        simple_frontier_max_fit_threshold=None,
    )
    assert calls


def test_split_mixed_order_has_single_type_capacities():
    ss = {"S": shipper()}
    points = generate_packings(
        ss, {"A": (1, 1, 1), "B": (1, 1, 1)}, ["A", "B"], {("S", "A"): 1, ("S", "B"): 1}, [[1, 1]]
    )
    total, usage = calculate_cost_count(
        [{"frequency": 3, "counts": {"A": 1, "B": 1}}], ss, ["A", "B"], points
    )
    assert (total, usage) == (6, {"S": 6})


def test_max_units_enforced_even_in_interpolation():
    points = comingle_frontier_points(
        ("S", 10, 1, 1, 3, 10),
        [("A", 1, 1, 1, 10), ("B", 1, 1, 1, 10)],
        [[1, 1]],
        shipper_constrained_by_max_units=True,
    )
    assert points and all(sum(p.values()) <= 3 for p in points)


def test_failed_profile_stops_analysis():
    with pytest.raises(AnalysisError) as caught:
        calculate_cost_count(
            [{"frequency": 4, "counts": {"A": 1}}, {"frequency": 6, "counts": {"B": 1}}],
            {"S": shipper()},
            ["A", "B"],
            [{"shipper_id": "S", "counts": {"A": 1}}],
        )
    assert caught.value.issues[0]["frequency"] == 6


def test_objective_choice_changes_solution():
    ss = {"SMALL": shipper(1, 10), "LARGE": shipper(2, 1)}
    pp = [{"shipper_id": sid, "counts": {"A": 1}} for sid in ss]
    profiles = [{"frequency": 1, "counts": {"A": 1}}]
    assert calculate_cost_count(profiles, ss, ["A"], pp)[1] == {"SMALL": 1, "LARGE": 0}
    assert calculate_cost_count(profiles, ss, ["A"], pp, objective="shipping_cost")[1] == {
        "SMALL": 0,
        "LARGE": 1,
    }


def test_seeded_preview_is_repeatable_without_extrapolation():
    profiles = [{"frequency": n, "counts": {"C": n}} for n in range(1, 8)] + [
        {"frequency": 0, "counts": {"C": 8}}
    ]
    args = (profiles, {"C": (1, 1, 1)}, {"S": shipper(10, 5)})
    a = analyze_loaded(*args, sample_ratio=0.5, random_seed=32)
    b = analyze_loaded(*args, sample_ratio=0.5, random_seed=32)
    assert a.to_dict() == b.to_dict()
    assert 0.5 <= a.coverage["included_fraction"] < 1
    assert a.total_cost == a.coverage["included_frequency"] * 5


@pytest.mark.parametrize("value", [-1, 1.5, "inf"])
def test_history_does_not_truncate_or_accept_bad_counts(tmp_path, value):
    p = tmp_path / "bad.xlsx"
    workbook(p, {"H": [["C", "FREQUENCY"], [value, 1]]})
    with pytest.raises(InputError) as caught:
        build_shipment_profiles(p, "H", ["C"])
    assert "C2" in caught.value.issues[0]["location"]


def test_duplicate_headers_and_uncached_formulas_rejected(tmp_path):
    p = tmp_path / "bad.xlsx"
    workbook(p, {"D": [["C", "C"], [1, 2]], "F": [["C", "FREQUENCY"], ["=1+1", 1]]})
    with pytest.raises(InputError, match="Duplicate"):
        read_xlsx_as_dicts(p, "D")
    with pytest.raises(InputError, match="formula"):
        build_shipment_profiles(p, "F", ["C"])


def test_empty_filtered_analysis_is_actionable():
    with pytest.raises(ValueError, match="No demand remains"):
        analyze_loaded(
            [{"frequency": 1, "counts": {"C": 1}}],
            {"C": (1, 1, 1)},
            {"S": shipper()},
            vector_freq_limit=100,
        )


@pytest.mark.parametrize("dims,cost", [((0, 1, 1), 0), ((1, 1, 1), -1)])
def test_bad_cost_inputs(dims, cost):
    with pytest.raises(ValueError):
        shipping_cost_zone4(dims, cost)


def test_sample_preparation_reconciles_ids_orders_units():
    config, base = load_config(ROOT / "examples/preparation.json")
    result = prepare_shipment_history(config, base)
    assert result["statistics"][0]["input_rows"] == 9
    assert result["statistics"][0]["included_orders"] == 8
    assert result["statistics"][0]["included_units"] == 18
    assert sum(p["frequency"] for p in result["sheets"]["East"]) == 8


def preparation_fixture(tmp_path):
    p = tmp_path / "raw.xlsx"
    workbook(
        p,
        {
            "H": [
                ["ORDER", "DEST", "SKU", "QTY"],
                ["001", "X", "001", 2],
                ["001", "X", "001", 2],
                ["001", "Y", "001", 1],
                ["002", "X", "001", 1],
                ["002", "X", "unknown", 1],
            ],
            "Map": [["SKU", "CARTON"], ["001", "C"]],
        },
    )
    return {
        "input_file": str(p),
        "sheets": ["H"],
        "mapping_file": str(p),
        "mapping_sheet": "Map",
        "sku_column": "SKU",
        "mapping_sku_column": "SKU",
        "mapping_carton_column": "CARTON",
        "order_key_columns": ["ORDER", "DEST"],
        "quantity_mode": "column",
        "quantity_column": "QTY",
    }


def test_unmatched_sku_stops_preparation(tmp_path):
    with pytest.raises(InputError, match="No partial profiles") as caught:
        prepare_shipment_history(preparation_fixture(tmp_path))
    assert "unknown" in caught.value.issues[0]["message"]


def test_explicit_exclusion_quarantines_whole_order_and_preserves_duplicates(tmp_path):
    config = preparation_fixture(tmp_path) | {"exclude_skus": ["unknown"]}
    result = prepare_shipment_history(config)
    stats = result["statistics"][0]
    assert stats["included_orders"] == 2
    assert stats["included_units"] == 5
    assert stats["excluded_orders"] == 1
    assert stats["duplicate_looking_rows_retained"] == 1
    one_per_row = prepare_shipment_history(config | {"quantity_mode": "one_per_row"})
    assert one_per_row["statistics"][0]["included_units"] == 3


def test_hand_calculated_comparison_and_workbook(tmp_path):
    config, base = load_config(ROOT / "examples/comparison.json")
    results, comparison, usage = run_scenarios(config, base)
    assert results[0].total_cost == 144
    assert results[1].total_cost == 120
    assert comparison[0]["cost_change"] == -24
    report = export_report(results, tmp_path, comparison, usage)
    wb = load_workbook(report, data_only=True)
    assert wb["Summary"]["B4"].value == 144
    assert wb["Summary"]["B5"].value == 120
    assert wb["Comparison"]["E4"].value == -24
    assert all(c.data_type != "e" for s in wb for row in s for c in row)
    wb.close()


def test_report_text_cannot_be_interpreted_as_formula(tmp_path):
    config, base = load_config(ROOT / "examples/scenario.json")
    result = run_scenario(config | {"name": "=1+1"}, base)
    wb = load_workbook(export_report([result], tmp_path), data_only=False)
    assert wb["Summary"]["A4"].data_type == "s"
    wb.close()


def test_unknown_config_option_rejected():
    config, base = load_config(ROOT / "examples/scenario.json")
    with pytest.raises(InputError, match="Unknown"):
        run_scenario(config | {"sample_raito": 0.1}, base)


def test_comparison_rejects_different_demand():
    config, base = load_config(ROOT / "examples/comparison.json")
    config["scenarios"][1]["history_sheet"] = "West"
    with pytest.raises(InputError, match="identical demand"):
        run_scenarios(config, base)


def test_cli_paths_with_spaces_and_error_report(tmp_path):
    config, base = load_config(ROOT / "examples/scenario.json")
    config["history_file"] = str(base / config["history_file"])
    config["reference_file"] = str(base / config["reference_file"])
    config["cartons"] = ["MISSING"]
    p = tmp_path / "bad scenario.json"
    p.write_text(json.dumps(config))
    out = tmp_path / "output directory"
    assert main(["run", str(p), "--output-dir", str(out)]) == 2
    assert len(list(out.glob("*/issues.json"))) == 1
    assert not list(out.glob("*/results.json"))


def test_cli_validation_has_no_output_on_success(tmp_path):
    assert (
        main(
            [
                "validate",
                str(ROOT / "examples/scenario.json"),
                "--output-dir",
                str(tmp_path / "unused"),
            ]
        )
        == 0
    )
    assert not (tmp_path / "unused").exists()


def test_selector_spread_matches_reference():
    points = [{"A": a, "B": 20 - a} for a in range(21)]
    result = _select_spread_packings(points, ["A", "B"], {"A": 20, "B": 20}, 5)
    assert result == [points[0], points[20], points[10], points[5], points[15]]


def test_preparation_preserves_distinct_numeric_text_skus(tmp_path):
    p = tmp_path / "text_ids.xlsx"
    workbook(
        p,
        {
            "H": [["ORDER", "SKU"], ["0001", "001"], ["0001", "1"]],
            "Map": [["SKU", "CARTON"], ["001", "C_A"], ["1", "C_B"]],
        },
    )
    config = {
        "input_file": str(p),
        "sheets": ["H"],
        "mapping_file": str(p),
        "mapping_sheet": "Map",
        "sku_column": "SKU",
        "mapping_sku_column": "SKU",
        "mapping_carton_column": "CARTON",
        "order_key_columns": ["ORDER"],
        "quantity_mode": "one_per_row",
    }
    profiles = prepare_shipment_history(config)["sheets"]["H"]
    assert profiles == [{"frequency": 1, "counts": {"C_A": 1, "C_B": 1}}]


def test_saved_configuration_resolves_inputs_from_output_folder(tmp_path):
    assert main(["run", str(ROOT / "examples/scenario.json"), "--output-dir", str(tmp_path)]) == 0
    snapshot = next(tmp_path.glob("*/configuration.json"))
    config = json.loads(snapshot.read_text())
    assert Path(config["history_file"]).is_absolute()
    assert main(["validate", str(snapshot)]) == 0


def test_unwritable_output_path_has_clean_failure(tmp_path, capsys):
    blocked = tmp_path / "is a file"
    blocked.write_text("existing content")
    assert main(["run", str(ROOT / "examples/scenario.json"), "--output-dir", str(blocked)]) == 2
    assert "Could not write the issue report" in capsys.readouterr().err
    assert blocked.read_text() == "existing content"


def test_input_changes_during_run_are_detected(tmp_path):
    config, base = load_config(ROOT / "examples/scenario.json")
    reference = tmp_path / "reference.xlsx"
    reference.write_bytes((base / config["reference_file"]).read_bytes())
    config["reference_file"] = str(reference)
    modified = False

    def progress(message):
        nonlocal modified
        if not modified:
            with reference.open("ab") as f:
                f.write(b"\n")
            modified = True

    with pytest.raises(InputError, match="input changed"):
        run_scenario(config, base, progress)


def test_legacy_and_saved_workflow_respect_the_same_selected_order():
    from cartonization.analysis import scenario_analysis

    config, base = load_config(ROOT / "examples/scenario.json")
    config.pop("shipper_overrides")
    config["cartons"] = ["C_B", "C_A"]
    config["shippers"] = ["S_SMALL", "S_BIG"]
    current = run_scenario(config, base)
    legacy = scenario_analysis(
        str(base / config["history_file"]),
        "East",
        "FREQUENCY",
        1,
        2,
        0,
        str(base / config["reference_file"]),
        ["LENGTH", "WIDTH", "HEIGHT"],
        "container_shipper_dims",
        "SHIPPER",
        "MAXUNITS",
        "container_carton_dims",
        "CARTON",
        config["cartons"],
        config["shippers"],
        shipping_cost_zone4,
        shipper_fixed_cost=2,
    )
    assert legacy[0] == current.total_cost
    assert legacy[1] == current.total_count
    assert legacy[2] == round(current.packing_efficiency)
    assert legacy[3] == current.packings
    assert legacy[4] == current.max_fits
    assert legacy[5] == current.shipment_profiles
